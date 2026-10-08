#!/usr/bin/env python3
"""cra/poam-to-xlsx.py — render the official-style POA&M workbook from poam.mjs JSON.

Reads reports/cra/poam/<product>.json (produced by `node cra/poam.mjs`) and writes
<product>.xlsx with the tabs a FedRAMP ConMon reviewer expects: Cover, Open POA&M Items,
Closed POA&M Items, Deviations. Data-export workbook (not a model): rows come from the JSON;
the Cover's summary counts are live COUNTIF/COUNTA formulas over the Open sheet so they stay
correct if rows are edited. Run scripts/recalc.py after (this script does it automatically).

  python3 cra/poam-to-xlsx.py [--product <id>] [--dir reports/cra/poam]

Requires openpyxl (preinstalled). Not the government template itself — a structured working
artifact teams reconcile into their official POA&M.
"""
import argparse, glob, json, os, subprocess, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3B57")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name=FONT, bold=True, size=15)
BODY = Font(name=FONT, size=10)
DUE_FILL = PatternFill("solid", fgColor="F4CCCC")      # past due
DEV_FILL = PatternFill("solid", fgColor="FFF2CC")      # deviation requested
KEV_FONT = Font(name=FONT, size=10, bold=True, color="B00020")
THIN = Border(*[Side(style="thin", color="D9DDE3")] * 4)


def _style_header(ws, ncol, row=1):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT; cell.border = THIN
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(ncol)}{row}"


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_rows(ws, rows, start=2):
    for r, row in enumerate(rows, start=start):
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = BODY; cell.border = THIN
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build(data, path):
    wb = Workbook()

    # ── Cover ────────────────────────────────────────────────────────────────
    cov = wb.active; cov.title = "Cover"
    cov["A1"] = "Plan of Action & Milestones (POA&M)"; cov["A1"].font = TITLE_FONT
    p = data["product"]; s = data["summary"]
    meta = [
        ("Product", f'{p["name"]} {p["version"]}  (id: {p["id"]})'),
        ("Generated", data["generatedAt"]),
        ("Evidence slice", data.get("slice") or "n/a"),
        ("Scope", "Code-side findings (dependencies / SAST / secrets) mapped to NIST SP 800-53 Rev 5."),
        ("Remediation SLA", "From DISCOVERY date — Critical/High 30 days, Moderate 90, Low 180 (FedRAMP)."),
        ("", ""),
        ("Open items", "=COUNTA('Open POA&M Items'!A2:A100000)"),
        ("  past due", '=COUNTIF(\'Open POA&M Items\'!O2:O100000,"*PAST DUE*")'),
        ("  deviation requested", '=COUNTIF(\'Open POA&M Items\'!O2:O100000,"*deviation*")'),
        ("Closed items (verified)", "=COUNTA('Closed POA&M Items'!A2:A100000)"),
        ("Deviations", "=COUNTA(Deviations!A2:A100000)"),
    ]
    for i, (k, v) in enumerate(meta, start=3):
        cov.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=True, size=10)
        cov.cell(row=i, column=2, value=v).font = BODY
    note = ("This is a generated working artifact, NOT the official FedRAMP POA&M template "
            "(proprietary formats cannot be submitted). Reconcile rows into your authorized "
            "template. Each row cites the tool + slice that found it (raw-scan → row traceability). "
            "Advisory control mapping — review with your assessor.")
    cov.cell(row=len(meta) + 4, column=1, value=note).font = Font(name=FONT, italic=True, size=9, color="666666")
    cov.merge_cells(start_row=len(meta) + 4, start_column=1, end_row=len(meta) + 4, end_column=2)
    cov.cell(row=len(meta) + 4, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    _widths(cov, [26, 96])

    # ── Open POA&M Items ─────────────────────────────────────────────────────
    ws = wb.create_sheet("Open POA&M Items")
    head = ["POA&M ID", "NIST 800-53 Controls", "Weakness", "Detection Source", "Source Identifier",
            "Asset Identifier", "Point of Contact", "Severity", "CVSS", "KEV", "EPSS",
            "Discovery Date", "SLA (days)", "Scheduled Completion", "Status", "Deviation", "Milestone / Remediation", "Slice"]
    ws.append(head); _style_header(ws, len(head))
    rows = [[r["poamId"], " ".join(r["controls"]), r["weaknessName"], r["weaknessSource"], r["sourceIdentifier"],
             r["asset"], r["pointOfContact"], r["severity"], r.get("cvss"), "YES" if r["kev"] else "",
             r.get("epss") if r.get("epss") is not None else "", r.get("discoveryDate") or "UNKNOWN",
             r["slaDays"], r.get("scheduledCompletionDate") or "n/a (no discovery date)", r["status"],
             (f'{r["deviation"]["kind"]}: {r["deviation"]["reason"]}' if r.get("deviation") else ""),
             r["milestone"], r.get("slice") or ""] for r in data["open"]]
    _write_rows(ws, rows)
    for i, r in enumerate(data["open"], start=2):
        if r["overdue"]:
            for c in range(1, len(head) + 1): ws.cell(row=i, column=c).fill = DUE_FILL
        elif r.get("deviation"):
            for c in range(1, len(head) + 1): ws.cell(row=i, column=c).fill = DEV_FILL
        if r["kev"]: ws.cell(row=i, column=10).font = KEV_FONT
    _widths(ws, [12, 16, 40, 12, 16, 34, 20, 9, 7, 6, 7, 20, 8, 20, 20, 26, 34, 20])

    # ── Closed POA&M Items ───────────────────────────────────────────────────
    wc = wb.create_sheet("Closed POA&M Items")
    head2 = ["POA&M ID", "NIST 800-53 Controls", "Weakness", "Source Identifier", "Asset Identifier",
             "Severity", "Discovery Date", "Remediation Date", "Change", "Evidence", "Status"]
    wc.append(head2); _style_header(wc, len(head2))
    _write_rows(wc, [[r["poamId"], " ".join(r["controls"]), r["weaknessName"], r["sourceIdentifier"], r["asset"],
                      r["severity"], r.get("discoveryDate") or "UNKNOWN", r.get("remediationDate") or "",
                      r["change"], r["evidence"], r["status"]] for r in data["closed"]])
    _widths(wc, [12, 16, 30, 16, 34, 9, 20, 20, 28, 50, 10])

    # ── Deviations ───────────────────────────────────────────────────────────
    wd = wb.create_sheet("Deviations")
    head3 = ["POA&M ID", "Source Identifier", "Asset", "Kind", "Type", "Reason", "Requested By", "At", "AO Approval Required"]
    wd.append(head3); _style_header(wd, len(head3))
    _write_rows(wd, [[d["poamId"], d["sourceIdentifier"], d["asset"], d["kind"], d["label"], d["reason"],
                      d["requestedBy"], d.get("at") or "", "YES" if d["needsAOApproval"] else "no"] for d in data["deviations"]])
    _widths(wd, [12, 16, 34, 8, 24, 40, 16, 20, 20])

    wb.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--product")
    ap.add_argument("--dir", default=os.path.join("reports", "cra", "poam"))
    a = ap.parse_args()
    files = [os.path.join(a.dir, f"{a.product}.json")] if a.product else sorted(glob.glob(os.path.join(a.dir, "*.json")))
    if not files or not os.path.exists(files[0]):
        sys.exit(f"no POA&M JSON in {a.dir} — run `node cra/poam.mjs` first")
    here = os.path.dirname(os.path.abspath(__file__))
    recalc = os.path.join(here, "..", "..", "..", ".claude", "skills", "xlsx", "scripts", "recalc.py")
    for f in files:
        data = json.load(open(f))
        out = f[:-5] + ".xlsx"
        build(data, out)
        if os.path.exists(recalc):
            r = subprocess.run([sys.executable, recalc, out], capture_output=True, text=True)
            try:
                st = json.loads(r.stdout).get("status")
            except Exception:
                st = r.stdout.strip()[:80]
            print(f"  ✓ {os.path.basename(out)}  (recalc: {st})")
        else:
            print(f"  ✓ {os.path.basename(out)}  (recalc skipped — scripts/recalc.py not found)")


if __name__ == "__main__":
    main()
